"""Customer branch XLSX UI checks. All API traffic is synthetic; no business writes."""
import argparse
import hashlib
import json
import tempfile
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook, Workbook
from playwright.sync_api import sync_playwright, expect

parser = argparse.ArgumentParser()
parser.add_argument('--url', default='http://127.0.0.1:3009')
parser.add_argument('--session-cookie-file', required=True)
parser.add_argument('--output', default='.backups/customer-store-ui')
args = parser.parse_args()
output = Path(args.output)
output.mkdir(parents=True, exist_ok=True)
checks, errors, blocked, requests = [], [], [], []
stores = [{'id': 23, 'name': 'Cơ sở QA A', 'code': 'QA-A'}, {'id': 31, 'name': 'Cơ sở QA B', 'code': 'QA-B'}]
customers = [
    {'id': 1, 'name': 'Khách QA một', 'id_card': '001234567890', 'store_id': None, 'phone': '0900000001', 'status': 'active'},
    {'id': 2, 'name': 'Khách QA hai', 'id_card': '001234567891', 'store_id': 23, 'phone': '0900000002', 'status': 'active'},
    {'id': 3, 'name': 'Khách QA ba', 'id_card': '001234567892', 'store_id': 31, 'phone': '0900000003', 'status': 'active'},
    {'id': 4, 'name': 'Khách QA bốn', 'id_card': '001234567893', 'store_id': None, 'phone': '0900000004', 'status': 'active'},
]
mode = {'next': ''}


def mock_api(route):
    request = route.request
    endpoint = urlparse(request.url).path
    if endpoint == '/api/auth/customers/match-stores' and request.method == 'POST':
        payload = request.post_data_json
        requests.append(payload)
        if mode['next'] == 'offline':
            mode['next'] = ''
            route.abort('failed')
            return
        if mode['next'] == 'conflict' and payload['commit']:
            mode['next'] = ''
            route.fulfill(status=409, json={'status': 'error', 'message': 'Khách hàng, cơ sở hoặc file đã thay đổi. Nhấn Kiểm tra lại trước khi cập nhật.'})
            return
        rows = []
        cards = [item['values']['id_card'].replace(' ', '') for item in payload['rows']]
        for item in payload['rows']:
            values = item['values']
            card = values['id_card'].replace(' ', '')
            customer = next((row for row in customers if row['id_card'] == card), None)
            branch = next((row for row in stores if values['store'] in [row['name'], row['code'], str(row['id'])]), None)
            previous = next((row for row in stores if customer and row['id'] == customer['store_id']), None)
            row_errors = list(item.get('errors', []))
            if not card.isdigit() or len(card) not in [9, 12]:
                row_errors.append('Căn cước sai định dạng.')
            if not customer:
                row_errors.append('Không tìm thấy khách hàng có căn cước này.')
            if not branch:
                row_errors.append('Cơ sở không tồn tại hoặc tên bị trùng.')
            if cards.count(card) > 1:
                row_errors.append('Căn cước trùng trong file.')
            state = 'invalid' if row_errors else 'unchanged' if customer['store_id'] == branch['id'] else 'ready'
            rows.append({'rowNumber': item['rowNumber'], 'values': values, 'customer_id': customer['id'] if customer else None,
                         'customer_name': customer['name'] if customer else '', 'previous_store_id': customer['store_id'] if customer else None,
                         'previous_store_name': previous['name'] if previous else '', 'store_id': branch['id'] if branch else None,
                         'store_name': branch['name'] if branch else '', 'state': state, 'errors': row_errors})
        revision = hashlib.sha256(json.dumps(rows, ensure_ascii=False).encode()).hexdigest()
        ready = sum(row['state'] == 'ready' for row in rows)
        if payload['commit']:
            assert payload['revision'] == revision, 'UI must submit the revision from its preview'
            assert ready > 0
            for row in rows:
                if row['state'] == 'ready':
                    next(customer for customer in customers if customer['id'] == row['customer_id'])['store_id'] = row['store_id']
        route.fulfill(json={'status': 'success', 'data': {'rows': rows, 'total': len(rows), 'ready': ready,
                      'unchanged': sum(row['state'] == 'unchanged' for row in rows), 'invalid': sum(row['state'] == 'invalid' for row in rows),
                      'updated': ready if payload['commit'] else 0, 'committed': payload['commit'], 'revision': revision}})
    elif request.method == 'GET':
        if endpoint == '/api/session':
            route.fulfill(json={'user': {'id': 143, 'name': 'QA'}})
            return
        data = customers if endpoint == '/api/auth/customers' else [{**row, 'store_name': row['name'], 'status': 1} for row in stores] if endpoint == '/api/auth/stores' else []
        route.fulfill(json={'status': 'success', 'data': data})
    else:
        blocked.append(endpoint)
        route.abort('blockedbyclient')


with tempfile.TemporaryDirectory(prefix='himoto-store-excel-ui-') as temp, sync_playwright() as playwright:
    temp = Path(temp)
    browser = playwright.chromium.launch()
    context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True, reduced_motion='reduce')
    cookie = Path(args.session_cookie_file).read_text(encoding='utf-8').strip()
    context.add_cookies([{'name': 'himoto_management_session', 'value': cookie, 'url': args.url, 'httpOnly': True, 'sameSite': 'Strict'}])
    context.route('**/api/**', mock_api)
    page = context.new_page()
    page.on('pageerror', lambda error: errors.append(str(error)))
    try:
        page.goto(args.url + '/customers', wait_until='domcontentloaded', timeout=120000)
        action = page.get_by_role('button', name='Khớp cơ sở theo căn cước', exact=True)
        expect(action).to_be_enabled(timeout=30000)
        page.add_style_tag(content='.mg-user-menu strong { visibility: hidden !important; }')
        action.click()
        dialog = page.get_by_role('dialog')
        with page.expect_download(timeout=30000) as event:
            dialog.get_by_role('button', name='Tải mẫu Excel căn cước / cơ sở', exact=True).click()
        template = temp / 'template.xlsx'
        event.value.save_as(template)
        book = load_workbook(template)
        assert list(book.worksheets[0].values)[0] == ('Căn cước', 'Cơ sở')
        assert book.worksheets[0]['A2'].number_format == '@' and book.worksheets[0]['A2'].value is None
        assert book['Cơ sở']['C2'].value == stores[0]['name']
        assert len(book.worksheets[0].data_validations.dataValidation) > 0
        checks.append('downloaded template has exactly two data columns, blank Text identities and current branch names/dropdown')

        def make_file(name, data):
            workbook = load_workbook(template)
            for i, row in enumerate(data, 2):
                for j, value in enumerate(row, 1):
                    workbook.worksheets[0].cell(i, j, value)
            filename = temp / name
            workbook.save(filename)
            return filename

        mixed = make_file('mixed.xlsx', [[customers[0]['id_card'], stores[0]['name']], [customers[1]['id_card'], stores[0]['name']],
                                        ['999999999999', stores[0]['name']], [customers[2]['id_card'], 'Không có cơ sở'],
                                        [customers[3]['id_card'], stores[0]['name']], [customers[3]['id_card'], stores[1]['name']]])
        file_input = dialog.get_by_label('Chọn file Excel căn cước / cơ sở (.xlsx)', exact=True)
        file_input.set_input_files(mixed)
        save = dialog.get_by_role('button', name='Cập nhật cơ sở cho 1 khách hàng', exact=True)
        expect(save).to_be_enabled(timeout=30000)
        expect(dialog.locator('tbody tr')).to_have_count(6)
        expect(dialog.locator('tbody tr').first).to_contain_text('Chưa có cơ sở')
        expect(dialog.locator('tbody tr').nth(1)).to_contain_text('Đã đúng cơ sở')
        assert requests[-1]['commit'] is False and requests[-1]['rows'][0]['values']['id_card'] == '001234567890'
        assert not any(request['commit'] for request in requests)
        dialog.get_by_label('Chỉ xem dòng lỗi', exact=True).check()
        expect(dialog.locator('tbody tr')).to_have_count(4)
        dialog.get_by_label('Chỉ xem dòng lỗi', exact=True).uncheck()
        page.screenshot(path=str(output / 'preview-1440.png'), full_page=True)
        checks.append('real XLSX upload matches identity/branch preview, keeps leading zeros, reports unknown/duplicate rows, distinguishes unchanged rows and writes only after confirmation')

        for width, height in [(375, 900), (844, 390)]:
            page.set_viewport_size({'width': width, 'height': height})
            page.wait_for_timeout(200)
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            bounds = dialog.bounding_box()
            assert bounds['x'] >= 0 and bounds['x'] + bounds['width'] <= width
            for _ in range(12):
                page.keyboard.press('Tab')
                assert dialog.evaluate('(node) => node.contains(document.activeElement)')
            page.screenshot(path=str(output / f'preview-{width}.png'), full_page=True)
        checks.append('375px and landscape layouts stay within viewport, table scrolls internally and keyboard focus remains inside dialog')

        save.click()
        expect(dialog.get_by_text('Đã cập nhật cơ sở cho 1 khách hàng. Giữ nguyên 1 dòng đã đúng; bỏ qua 4 dòng lỗi.', exact=True)).to_be_visible(timeout=30000)
        assert requests[-1]['commit'] is True and len(requests[-1]['rows']) == 6
        assert customers[0]['store_id'] == 23 and customers[3]['store_id'] is None
        expect(file_input).to_be_disabled()
        dialog.get_by_role('button', name='Đóng', exact=True).click()
        expect(page.locator('tbody tr').filter(has_text='Khách QA một')).to_contain_text(stores[0]['name'])
        checks.append('confirmed update sends complete preview plus revision, changes only eligible branch, reports saved/skipped counts and refreshes customer list')

        page.set_viewport_size({'width': 1440, 'height': 1000})
        action.click()
        wrong = temp / 'wrong.xls'
        wrong.write_text('wrong')
        file_input.set_input_files(wrong)
        expect(dialog.get_by_role('alert')).to_contain_text('Chỉ nhận file Excel .xlsx')
        empty = make_file('empty.xlsx', [])
        file_input.set_input_files(empty)
        expect(dialog.get_by_role('alert')).to_contain_text('File chưa có dữ liệu')
        numeric = make_file('numeric.xlsx', [[1234567890, stores[0]['name']]])
        file_input.set_input_files(numeric)
        expect(dialog.locator('.mg-import-result')).to_contain_text('phải là Text')
        expect(dialog.get_by_role('button', name='Cập nhật cơ sở cho 0 khách hàng', exact=True)).to_be_disabled()
        checks.append('unsupported/empty files and numeric identities show explicit errors and disable saving')

        retry = temp / 'reordered.xlsx'
        reversed_book = Workbook()
        reversed_book.active.append(['Tên cơ sở', 'CCCD'])
        reversed_book.active.append([stores[0]['name'], customers[2]['id_card']])
        reversed_book.save(retry)
        mode['next'] = 'offline'
        file_input.set_input_files(retry)
        expect(dialog.get_by_role('alert')).to_be_visible()
        dialog.get_by_role('button', name='Kiểm tra lại', exact=True).click()
        expect(save).to_be_enabled()
        assert requests[-1]['rows'][0]['values']['store'] == stores[0]['name']
        mode['next'] = 'conflict'
        save.click()
        expect(dialog.get_by_role('alert')).to_contain_text('đã thay đổi')
        expect(dialog.get_by_role('button', name='Cập nhật cơ sở cho 0 khách hàng', exact=True)).to_be_disabled()
        dialog.get_by_role('button', name='Kiểm tra lại', exact=True).click()
        expect(save).to_be_enabled()
        save.click()
        expect(dialog.get_by_text('Đã cập nhật cơ sở cho 1 khách hàng. Giữ nguyên 0 dòng đã đúng; bỏ qua 0 dòng lỗi.', exact=True)).to_be_visible()
        checks.append('reordered aliases parse correctly; network retry retains inputs; stale preview requires recheck before another save')
        assert not errors, errors
        assert not blocked, blocked
    finally:
        report = {'passed': len(checks), 'checks': checks, 'javascriptErrors': errors, 'unexpectedWritesBlocked': blocked, 'realBusinessWrites': 0}
        (output / 'results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(report, ensure_ascii=False))
        browser.close()
