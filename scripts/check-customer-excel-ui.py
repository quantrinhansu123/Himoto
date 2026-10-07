"""Excel UI checks with synthetic API responses. Never sends a business write to DB."""
import argparse
import json
import tempfile
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook, Workbook
from playwright.sync_api import sync_playwright, expect

parser = argparse.ArgumentParser()
parser.add_argument('--url', default='http://127.0.0.1:3004')
parser.add_argument('--session-cookie-file', required=True)
parser.add_argument('--output', default='docs/qa/customer-excel')
args = parser.parse_args()
output = Path(args.output)
output.mkdir(parents=True, exist_ok=True)
checks, errors, blocked, requests = [], [], [], []
customers = [{'id': 99, 'name': 'Khách đã có QA', 'phone': '0900000099', 'id_card': '009999999999', 'status': 'active', 'store_id': 2, 'address': 'Địa chỉ QA'}]
mode = {'next': ''}

def mock_api(route):
    request = route.request
    endpoint = urlparse(request.url).path
    if endpoint == '/api/auth/customers/import' and request.method == 'POST':
        payload = request.post_data_json
        requests.append(payload)
        if mode['next'] == 'offline':
            mode['next'] = ''
            route.abort('failed')
            return
        if mode['next'] == 'conflict' and payload['commit']:
            mode['next'] = ''
            route.fulfill(status=409, json={'status': 'error', 'message': 'Dữ liệu đã thay đổi. Kiểm tra lại trước khi nhập.'})
            return
        rows = []
        for item in payload['rows']:
            values = dict(item['values'])
            values['status'] = values['status'] or 'active'
            if values['status'] == 'Chưa hoàn tất':
                values['status'] = 'draft'
            row_errors = list(item.get('errors', []))
            warnings = []
            incomplete_allowed = payload.get('allowIncomplete') is True and values['status'] == 'draft'
            for field in ['name', 'phone', 'id_card', 'address', 'store']:
                if not values[field]:
                    if incomplete_allowed and field in ['id_card', 'address']:
                        warnings.append('Chưa có ' + ('cccd / cmnd' if field == 'id_card' else 'địa chỉ') + '; cần bổ sung hồ sơ sau khi nhập.')
                    else:
                        row_errors.append('Thiếu ' + field)
            if values['id_card'] and (not values['id_card'].isdigit() or len(values['id_card']) not in [9, 12]):
                row_errors.append('CCCD/CMND sai định dạng.')
            if values['email'] and '@' not in values['email']:
                row_errors.append('Email chưa đúng định dạng.')
            has_errors = bool(row_errors)
            duplicates = [record for record in customers if (values['id_card'] and record['id_card'] == values['id_card']) or (values['phone'] and record['phone'] == values['phone'])]
            if duplicates:
                row_errors.append('CCCD/CMND đã có ở khách hàng #99.')
            state = 'invalid' if has_errors else 'duplicate' if duplicates else 'valid'
            rows.append({'rowNumber': item['rowNumber'], 'values': values, 'store_id': 2, 'state': state, 'errors': row_errors, 'warnings': warnings})
        valid = sum(row['state'] == 'valid' for row in rows)
        result = {'rows': rows, 'total': len(rows), 'valid': valid, 'invalid': sum(row['state'] == 'invalid' for row in rows), 'duplicate': sum(row['state'] == 'duplicate' for row in rows), 'incomplete': sum(row['state'] == 'valid' and bool(row['warnings']) for row in rows), 'imported': valid if payload['commit'] else 0, 'committed': payload['commit']}
        if payload['commit']:
            assert valid == len(rows), 'UI must submit only previewed valid rows'
            for row in rows:
                customers.append({**row['values'], 'id': len(customers) + 100, 'store_id': 2})
        route.fulfill(status=201 if payload['commit'] else 200, json={'status': 'success', 'data': result})
    elif request.method == 'GET':
        if endpoint == '/api/session':
            route.fulfill(json={'user': {'id': 143, 'name': 'QA'}})
            return
        data = customers if endpoint == '/api/auth/customers' else [{'id': 2, 'code': 'CS-QA', 'store_name': 'Cơ sở QA', 'status': 1}] if endpoint == '/api/auth/stores' else []
        route.fulfill(json={'status': 'success', 'data': data})
    else:
        blocked.append(endpoint)
        route.abort('blockedbyclient')

with tempfile.TemporaryDirectory(prefix='himoto-excel-ui-') as temp, sync_playwright() as playwright:
    temp = Path(temp)
    browser = playwright.chromium.launch()
    context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True, reduced_motion='reduce')
    cookie = Path(args.session_cookie_file).read_text(encoding='utf-8').strip()
    context.add_cookies([{'name': 'himoto_management_session', 'value': cookie, 'url': args.url, 'httpOnly': True, 'sameSite': 'Strict'}])
    context.route('**/api/**', mock_api)
    page = context.new_page()
    page.on('pageerror', lambda error: errors.append(str(error)))
    try:
        page.goto(args.url + '/customers', wait_until='domcontentloaded', timeout=120000)
        expect(page.get_by_role('button', name='Tải mẫu Excel', exact=True)).to_be_enabled(timeout=30000)
        page.add_style_tag(content='.mg-user-menu strong { visibility: hidden !important; }')
        with page.expect_download(timeout=30000) as event:
            page.get_by_role('button', name='Tải mẫu Excel', exact=True).click()
        template = temp / 'template.xlsx'
        event.value.save_as(template)
        book = load_workbook(template)
        assert book.sheetnames == ['Khách hàng', 'Cơ sở', 'Hướng dẫn']
        assert book['Khách hàng']['B2'].number_format == '@'
        assert book['Cơ sở']['C2'].value == 'Cơ sở QA'
        assert book['Khách hàng']['A2'].value is None
        checks.append('browser download yields a valid blank XLSX with Text identity columns and current branch dropdown')

        def make_file(name, data):
            workbook = load_workbook(template)
            for index, row in enumerate(data, 2):
                for column, value in enumerate(row, 1):
                    workbook['Khách hàng'].cell(index, column, value)
            filename = temp / name
            workbook.save(filename)
            return filename

        good = ['Khách Excel QA', '0900000001', '001234567890', 'qa@example.invalid', 'Địa chỉ Excel QA', 'Cơ sở QA', '', '']
        invalid = ['Khách lỗi QA', '0900000002', '001234567891', 'bad-email', 'Địa chỉ QA', 'Cơ sở QA', '', '']
        duplicate = ['Khách trùng QA', '0900000099', '009999999999', '', 'Địa chỉ QA', 'Cơ sở QA', '', '']
        mixed = make_file('mixed.xlsx', [good, invalid, duplicate])
        page.get_by_role('button', name='Nhập Excel', exact=True).click()
        dialog = page.get_by_role('dialog')
        dialog.get_by_label('Chọn file Excel khách hàng (.xlsx)', exact=True).set_input_files(mixed)
        expect(dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True)).to_be_enabled(timeout=30000)
        assert requests[-1]['commit'] is False and len(requests[-1]['rows']) == 3
        assert requests[-1]['rows'][0]['values']['id_card'] == good[2]
        assert requests[-1]['rows'][0]['values']['phone'] == good[1]
        expect(dialog.locator('tbody tr')).to_have_count(3)
        dialog.get_by_label('Chỉ xem dòng lỗi / trùng', exact=True).check()
        expect(dialog.locator('tbody tr')).to_have_count(2)
        dialog.get_by_label('Chỉ xem dòng lỗi / trùng', exact=True).uncheck()
        page.screenshot(path=str(output / 'preview-1440.png'), full_page=True)
        checks.append('upload reads real XLSX, preserves leading zeros, previews valid/error/duplicate rows and filters issues')

        page.set_viewport_size({'width': 375, 'height': 900})
        page.wait_for_function("getComputedStyle(document.querySelector('.mg-customer-import')).maxWidth === '343px'")
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        bounds = dialog.bounding_box()
        assert bounds['x'] >= 0 and bounds['x'] + bounds['width'] <= 375, {'bounds': bounds, 'computed': dialog.evaluate('(node) => ({width: getComputedStyle(node).width, maxWidth: getComputedStyle(node).maxWidth, minWidth: getComputedStyle(node).minWidth, innerWidth, viewport: visualViewport.width})')}
        for _ in range(16):
            page.keyboard.press('Tab')
            assert dialog.evaluate('(node) => node.contains(document.activeElement)')
        page.screenshot(path=str(output / 'preview-375.png'), full_page=True)
        checks.append('375px preview has no page overflow, table scrolls internally and keyboard focus stays in dialog')

        dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True).click()
        expect(dialog.get_by_text('Đã nhập 1 khách hàng. Bỏ qua 1 dòng lỗi và 1 dòng trùng.', exact=True)).to_be_visible(timeout=30000)
        assert requests[-1]['commit'] is True and len(requests[-1]['rows']) == 1
        dialog.get_by_role('button', name='Đóng', exact=True).click()
        expect(page.locator('.mg-title-count')).to_have_text('2')
        checks.append('confirmation sends only eligible rows, shows saved/skipped counts and refreshes customer list (mock write)')

        page.get_by_role('button', name='Nhập Excel', exact=True).click()
        wrong = temp / 'wrong.xls'
        wrong.write_text('unsupported')
        dialog.get_by_label('Chọn file Excel khách hàng (.xlsx)', exact=True).set_input_files(wrong)
        expect(dialog.get_by_role('alert')).to_contain_text('Chỉ nhận file Excel .xlsx')
        expect(dialog.get_by_role('button', name='Nhập 0 khách hàng hợp lệ', exact=True)).to_be_disabled()
        numeric = make_file('numeric.xlsx', [[good[0], 900000001, 1234567890, *good[3:]]])
        dialog.get_by_label('Chọn file Excel khách hàng (.xlsx)', exact=True).set_input_files(numeric)
        expect(dialog.get_by_role('button', name='Kiểm tra lại', exact=True)).to_be_enabled()
        expect(dialog.locator('.mg-import-result')).to_contain_text('phải là Text')
        expect(dialog.get_by_role('button', name='Nhập 0 khách hàng hợp lệ', exact=True)).to_be_disabled()
        empty = make_file('blank.xlsx', [])
        dialog.get_by_label('Chọn file Excel khách hàng (.xlsx)', exact=True).set_input_files(empty)
        expect(dialog.get_by_role('alert')).to_contain_text('File chưa có khách hàng')
        checks.append('unsupported/blank files and numeric identities display errors and disable import')

        retry_good = good.copy()
        retry_good[1], retry_good[2] = '0900000003', '001234567893'
        retry = make_file('retry.xlsx', [retry_good])
        mode['next'] = 'offline'
        dialog.get_by_label('Chọn file Excel khách hàng (.xlsx)', exact=True).set_input_files(retry)
        expect(dialog.get_by_role('alert')).to_be_visible()
        expect(dialog.get_by_role('button', name='Kiểm tra lại', exact=True)).to_be_enabled()
        dialog.get_by_role('button', name='Kiểm tra lại', exact=True).click()
        expect(dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True)).to_be_enabled()
        mode['next'] = 'conflict'
        dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True).click()
        expect(dialog.get_by_role('alert')).to_contain_text('Dữ liệu đã thay đổi')
        expect(dialog.get_by_role('button', name='Nhập 0 khách hàng hợp lệ', exact=True)).to_be_disabled()
        dialog.get_by_role('button', name='Kiểm tra lại', exact=True).click()
        expect(dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True)).to_be_enabled()
        dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True).click()
        expect(dialog.get_by_text('Đã nhập 1 khách hàng. Bỏ qua 0 dòng lỗi và 0 dòng trùng.', exact=True)).to_be_visible()
        checks.append('network failure retains parsed data for retry; commit conflict forces new preview before another import')
        dialog.get_by_role('button', name='Đóng', exact=True).click()
        page.set_viewport_size({'width': 1440, 'height': 1000})
        page.get_by_role('button', name='Nhập Excel', exact=True).click()
        partial = ['Khách thiếu hồ sơ QA', '0900000008', '', '', '', 'Cơ sở QA', 'Chưa hoàn tất', 'Bổ sung hồ sơ QA']
        active_partial = ['Khách thiếu hồ sơ active QA', '0900000009', '', '', '', 'Cơ sở QA', 'Bình thường', '']
        wrong_card = ['Khách sai giấy tờ QA', '0900000010', '1234', '', '', 'Cơ sở QA', 'Chưa hoàn tất', '']
        partial_file = make_file('partial.xlsx', [partial, active_partial, wrong_card])
        dialog.get_by_label('Chọn file Excel khách hàng (.xlsx)', exact=True).set_input_files(partial_file)
        expect(dialog.get_by_role('button', name='Kiểm tra lại', exact=True)).to_be_enabled()
        expect(dialog.get_by_role('button', name='Nhập 0 khách hàng hợp lệ', exact=True)).to_be_disabled()
        assert requests[-1]['allowIncomplete'] is False
        dialog.get_by_label('Cho phép hồ sơ Chưa hoàn tất thiếu CCCD/địa chỉ', exact=True).check()
        expect(dialog.locator('tbody tr')).to_have_count(0)
        expect(dialog.get_by_role('button', name='Nhập 0 khách hàng hợp lệ', exact=True)).to_be_disabled()
        dialog.get_by_role('button', name='Kiểm tra lại', exact=True).click()
        expect(dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True)).to_be_enabled()
        assert requests[-1]['allowIncomplete'] is True and requests[-1]['commit'] is False
        expect(dialog.locator('tbody tr').first).to_contain_text('Chưa có cccd / cmnd')
        expect(dialog.locator('tbody tr').first).to_contain_text('Chưa có địa chỉ')
        page.screenshot(path=str(output / 'incomplete-1440.png'), full_page=True)
        page.set_viewport_size({'width': 375, 'height': 900})
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        page.screenshot(path=str(output / 'incomplete-375.png'), full_page=True)
        dialog.get_by_role('button', name='Nhập 1 khách hàng hợp lệ', exact=True).click()
        expect(dialog.get_by_text('Đã nhập 1 khách hàng. Bỏ qua 2 dòng lỗi và 0 dòng trùng.', exact=True)).to_be_visible()
        assert requests[-1]['allowIncomplete'] is True and requests[-1]['commit'] is True and len(requests[-1]['rows']) == 1
        assert customers[-1]['status'] == 'draft' and not customers[-1]['id_card'] and not customers[-1]['address']
        expect(dialog.get_by_label('Cho phép hồ sơ Chưa hoàn tất thiếu CCCD/địa chỉ', exact=True)).to_be_disabled()
        dialog.get_by_role('button', name='Đóng', exact=True).click()
        expect(page.locator('.mg-title-count')).to_have_text('4')
        expect(page.get_by_text('Khách thiếu hồ sơ QA', exact=True)).to_be_visible()
        expect(page.locator('tbody tr').filter(has_text='Khách thiếu hồ sơ QA').get_by_text('Chưa hoàn tất', exact=True)).to_be_visible()
        checks.append('incomplete mode requires opt-in and fresh preview, warns about missing documents, excludes active/malformed-ID rows, commits draft with notes and renders it at 1440/375px')
        assert not errors, errors
        assert not blocked, blocked
    finally:
        report = {'passed': len(checks), 'checks': checks, 'javascriptErrors': errors, 'unexpectedWritesBlocked': blocked, 'realBusinessWrites': 0, 'dataSource': 'synthetic API responses; server layout authentication reads existing QA account only'}
        (output / 'results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(report, ensure_ascii=False))
        browser.close()
